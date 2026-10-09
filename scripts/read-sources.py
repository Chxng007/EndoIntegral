import sys, json, shutil
from pathlib import Path
sys.path.insert(0, str(Path('.tools/python').resolve()))
import fitz
from docx import Document
root = Path('C:/Users/innovacion/Downloads')
out = Path('docs/sources')
out.mkdir(parents=True, exist_ok=True)
renders = Path('artifacts/references')
renders.mkdir(parents=True, exist_ok=True)
files = ['Meditaciones_Mindfulness_Guiadas_CLICABLE-2.pdf','Yoga_Terapia_Endometriosis_Posturas_y_Videos.pdf','página web.pdf','CONSIDERACIONES-LABORALES-FRENTE-A-TRABAJADORAS-DIAGNOSTICADAS-CON-ENDOMETRIOSIS.pdf']
inventory = []
for i, name in enumerate(files):
    doc = fitz.open(root / name)
    text = []
    for n, page in enumerate(doc):
        text.append(f'\n--- PÁGINA {n+1} ---\n' + page.get_text(sort=True))
        links = [l['uri'] for l in page.get_links() if 'uri' in l]
        text.extend(links)
        if i in [1, 2]:
            page.get_pixmap(matrix=fitz.Matrix(1.2,1.2)).save(renders / f'doc-{i}-page-{n+1}.png')
    (out / f'{i+1}-{Path(name).stem}.txt').write_text('\n'.join(text), encoding='utf-8')
    inventory.append({'name':name, 'pages':len(doc), 'characters':sum(map(len,text))})
doc = Document(root / 'Estructura de la Página Web y Contenido.docx')
parts = [p.text for p in doc.paragraphs]
for table in doc.tables:
    for row in table.rows:
        parts.append(' | '.join(c.text for c in row.cells))
(out / 'estructura.txt').write_text('\n'.join(parts), encoding='utf-8')
shutil.copyfile(root/'PROMPT_MAESTRO_EndoIntegral.md',out/'PROMPT_MAESTRO_EndoIntegral.md')
# Las guías son contenido de los planes: van al bucket privado «resources» de Supabase
# (ver 202610090012_protected_documents.sql), nunca a public/.
Path('docs/originales/protegidos/yoga').mkdir(parents=True,exist_ok=True)
Path('docs/originales/protegidos/mindfulness').mkdir(parents=True,exist_ok=True)
Path('public/img').mkdir(parents=True,exist_ok=True)
shutil.copyfile(root/files[0], 'docs/originales/protegidos/mindfulness/meditaciones-mindfulness.pdf')
shutil.copyfile(root/files[1], 'docs/originales/protegidos/yoga/yoga-terapia-endometriosis.pdf')
shutil.copyfile(root/'arf.jpeg', 'docs/sources/referencia-anatomica.jpeg')
print(json.dumps(inventory, ensure_ascii=False, indent=2))
